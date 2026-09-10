"""Tabular profiling for coverage and staleness — CSV, TSV, and XLSX."""

from __future__ import annotations

import csv
import hashlib
import re
import zipfile
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import Any

TABULAR_SUFFIXES = {".csv", ".tsv", ".xlsx"}


def profile_file(path: Path) -> dict[str, Any] | None:
    """Return a facts-shaped dict compatible with `staleness.extract_measures`, or None."""
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        return _profile_xlsx(path)
    if suffix in {".csv", ".tsv"}:
        return _profile_csv(path, delimiter="\t" if suffix == ".tsv" else ",")
    return None


def file_sha256(path: Path) -> str:
    """Hash file contents, or empty string on failure."""
    try:
        digest = hashlib.sha256()
        with path.open("rb") as handle:
            for chunk in iter(lambda: handle.read(65536), b""):
                digest.update(chunk)
        return digest.hexdigest()
    except OSError:
        return ""


def _profile_csv(path: Path, *, delimiter: str) -> dict[str, Any] | None:
    try:
        with path.open(newline="", encoding="utf-8") as handle:
            reader = csv.DictReader(handle, delimiter=delimiter)
            fieldnames = reader.fieldnames or []
            rows = list(reader)
    except (OSError, UnicodeDecodeError, csv.Error):
        return None

    if not fieldnames:
        return {"rows": 0, "numeric": []}

    return _facts_from_rows(list(fieldnames), rows)


def _profile_xlsx(path: Path) -> dict[str, Any] | None:
    try:
        from openpyxl import load_workbook
        from openpyxl.utils.exceptions import InvalidFileException
    except ImportError:
        return None

    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        if sheet is None:
            workbook.close()
            return {"rows": 0, "numeric": []}
        rows_iter = sheet.iter_rows(values_only=True)
        header = next(rows_iter, None)
        if not header:
            workbook.close()
            return {"rows": 0, "numeric": []}
        fieldnames = [str(cell) if cell is not None else "" for cell in header]
        rows = [
            {fieldnames[i]: ("" if value is None else value) for i, value in enumerate(row)}
            for row in rows_iter
        ]
        workbook.close()
    except (OSError, ValueError, KeyError, zipfile.BadZipFile, InvalidFileException):
        # A .xlsx is a zip archive, and a file named like one need not be. A CSV renamed by hand, a
        # truncated download, an HTML error page saved with the wrong extension: openpyxl raises
        # BadZipFile for those, InvalidFileException for a legacy .xls, and KeyError when the
        # archive opens but has no `xl/workbook.xml`. None of them was caught, so `assurance check`
        # died with a traceback mid-folder — with --json there was no JSON on stdout at all, and
        # through the MCP tool an agent got "Error executing tool" and nothing else.
        #
        # Returning None is the caller's existing "could not be read as a table" path: the period is
        # reported as unreadable, the rest of the folder is still checked, exit codes as usual.
        return None

    return _facts_from_rows(fieldnames, rows)


# --- what the file says about its own periods ------------------------------------------------
#
# The filename is a CLAIM about what a file contains, and until 0.5.9 nothing tested it. `check`
# opened every file, read every row into memory, kept the row count and the numeric totals, and
# threw the dates away — so a folder where January and February had been merged into one file
# reported February missing and failed the build over a month that was sitting right there, and a
# file named `2024-03.csv` holding February rows reported "3 of 3 months, complete".
#
# The read is already paid for. These functions keep what was being discarded.

_MONTHS = {
    m: i
    for i, name in enumerate(
        ("january", "february", "march", "april", "may", "june",
         "july", "august", "september", "october", "november", "december"),
        start=1,
    )
    for m in (name, name[:3])
}

# Unambiguous shapes only. `05/01/2024` is deliberately absent: it is the 5th of January to half
# the world and the 1st of May to the other half, and a warning built on a coin flip is worse than
# no warning. A column written that way is reported as unreadable rather than guessed at.
_ISO = re.compile(r"^(\d{4})[-/](\d{1,2})(?:[-/](\d{1,2}))?(?:[T ].*)?$")
_DMY_NAMED = re.compile(r"^(\d{1,2})[ \-]([A-Za-z]{3,9})[ \-](\d{4})$")
_MY_NAMED = re.compile(r"^([A-Za-z]{3,9})[ \-,]+(\d{4})$")

#: Rows scanned for dates before giving up. The whole file is already in memory; this bounds the
#: date PARSING, which is the part that costs on a very large corpus.
MAX_DATE_ROWS = 200_000

#: A column is a date column when this much of its non-empty values parse. Below it, the column is
#: something else that happens to contain a few dates.
_DATE_COLUMN_RATIO = 0.8


def parse_content_date(value: Any) -> date | None:
    """One cell into a date, or None when it is not unambiguously one."""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raw = str(value or "").strip()
    if not raw:
        return None
    match = _ISO.match(raw)
    if match:
        year, month, day = int(match.group(1)), int(match.group(2)), int(match.group(3) or 1)
        try:
            return date(year, month, day)
        except ValueError:
            return None
    match = _DMY_NAMED.match(raw)
    if match:
        # A separate name from the numeric `month` above: that one is an int, this one is an
        # int-or-None lookup, and reusing the name made the type of the variable depend on which
        # branch you arrived from.
        named = _MONTHS.get(match.group(2).lower())
        if named is None:
            return None
        try:
            return date(int(match.group(3)), named, int(match.group(1)))
        except ValueError:
            return None
    match = _MY_NAMED.match(raw)
    if match:
        named = _MONTHS.get(match.group(1).lower())
        if named is not None:
            return date(int(match.group(2)), named, 1)
    return None


def _dates_from_rows(fieldnames: list[str], rows: list[dict[str, Any]]) -> dict[str, Any]:
    """Every column that reads as dates, and which dates it holds.

    Deliberately reports ALL candidate columns rather than picking one. `created_at` and
    `report_date` are both dates and they are not both the period; choosing between them here would
    be a guess made where the caller cannot see it. The caller knows the series kind and can ask
    whether the candidates agree under it.
    """
    scanned = rows[:MAX_DATE_ROWS]
    columns: dict[str, dict[str, int]] = {}
    for column in fieldnames:
        seen = 0
        parsed: Counter[str] = Counter()
        for row in scanned:
            raw = row.get(column)
            if raw is None or str(raw).strip() == "":
                continue
            seen += 1
            found = parse_content_date(raw)
            if found is not None:
                parsed[found.isoformat()] += 1
        if seen and sum(parsed.values()) >= max(1, int(seen * _DATE_COLUMN_RATIO)):
            columns[column] = dict(parsed)
    return {
        "columns": columns,
        "rows_scanned": len(scanned),
        "truncated": len(rows) > len(scanned),
    }


def _facts_from_rows(fieldnames: list[str], rows: list[dict[str, Any]]) -> dict[str, Any]:
    numeric: list[dict[str, object]] = []
    row_count = len(rows)
    for column in fieldnames:
        values: list[float] = []
        for row in rows:
            raw = str(row.get(column) or "").strip()
            if not raw:
                continue
            cleaned = raw.replace(",", "").replace("$", "").replace("£", "").replace("€", "")
            if cleaned.endswith("%"):
                cleaned = cleaned[:-1]
            try:
                values.append(float(cleaned))
            except ValueError:
                values = []
                break
        if values and len(values) >= max(1, int(row_count * 0.8)):
            numeric.append({"name": column, "total": sum(values)})

    return {"rows": row_count, "numeric": numeric, "dates": _dates_from_rows(fieldnames, rows)}
